"""
Turns a bank or credit-card statement file (PDF, Excel, or an "xls" that is
really HTML or tab-separated text) into a plain grid of strings.

This module deliberately stops at the grid. Finding the header row, guessing
which column is the date or the debit, reading Indian amounts and dates, and
deduplicating all happen downstream in the frontend and backend, which already
do that for CSV uploads. Every file type meets the same pipeline once it is a
grid.

Nothing here keeps the file or the password: both live only for the duration
of one call.
"""
from __future__ import annotations

import csv
import io
import re
import statistics
from dataclasses import dataclass
from datetime import date, datetime

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 60
MAX_ROWS = 10_000
MAX_COLS = 30


class StatementError(Exception):
    """A statement we can't read, with a code the UI can act on."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class Extracted:
    grid: list[list[str]]
    format: str
    pages: int = 1


def extract(data: bytes, password: str | None = None) -> Extracted:
    """Detects the real file type from its bytes (not its name) and extracts a grid."""
    if not data:
        raise StatementError("empty", "The file is empty.")
    if len(data) > MAX_BYTES:
        raise StatementError("too_large", "The file is larger than 10 MB.")

    kind = sniff(data)
    if kind == "pdf":
        result = _extract_pdf(data, password)
    elif kind == "xlsx":
        result = Extracted(_extract_xlsx(data), "xlsx")
    elif kind == "xls":
        result = Extracted(_extract_xls(data), "xls")
    elif kind == "html":
        result = Extracted(_extract_html(data), "html")
    elif kind == "text":
        result = Extracted(_extract_text(data), "text")
    else:
        raise StatementError(
            "unsupported",
            "This doesn't look like a statement file. Upload the PDF, Excel or CSV "
            "your bank provides.")

    grid = _clean(result.grid)
    if len(grid) < 2:
        raise StatementError("no_table", "No transaction table was found in this file.")
    result.grid = grid
    return result


def sniff(data: bytes) -> str:
    head = data[:2048].lstrip(b"\xef\xbb\xbf \t\r\n")
    if data.startswith(b"%PDF") or b"%PDF-" in data[:1024]:
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        return "xlsx"
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        # OLE2: a real .xls — or a password-protected .xlsx, which Office
        # wraps in the same container
        if b"E\x00n\x00c\x00r\x00y\x00p\x00t\x00e\x00d\x00P\x00a\x00c\x00k\x00a\x00g\x00e" in data:
            raise StatementError(
                "excel_encrypted",
                "This Excel file is password-protected. Open it, remove the password "
                "(or save a copy as CSV), and upload that.")
        return "xls"
    lower = head.lower()
    if lower.startswith(b"<") and b"<table" in data[:200_000].lower():
        return "html"
    if _looks_like_text(data[:4096]):
        return "text"
    return "unknown"


def _looks_like_text(sample: bytes) -> bool:
    if b"\x00" in sample:
        return False
    try:
        sample.decode("utf-8")
        return True
    except UnicodeDecodeError:
        # Latin-1 exports from older core-banking systems
        printable = sum(32 <= b < 127 or b in (9, 10, 13) for b in sample)
        return printable / max(1, len(sample)) > 0.95


# ── PDF ─────────────────────────────────────────────────────────────────────

def _extract_pdf(data: bytes, password: str | None) -> Extracted:
    import pdfplumber
    from pdfminer.pdfdocument import PDFPasswordIncorrect
    from pdfplumber.utils.exceptions import PdfminerException

    try:
        pdf = pdfplumber.open(io.BytesIO(data), password=password or "")
    except PdfminerException as e:
        if any(isinstance(a, PDFPasswordIncorrect) for a in e.args):
            if password:
                raise StatementError("password_incorrect",
                                     "That password didn't open the PDF. Try again.") from None
            raise StatementError("password_required",
                                 "This PDF is password-protected.") from None
        raise StatementError("unreadable", "This PDF could not be read.") from None

    with pdf:
        if len(pdf.pages) > MAX_PAGES:
            raise StatementError(
                "too_many_pages",
                f"This PDF has {len(pdf.pages)} pages; the limit is {MAX_PAGES}. "
                "Download a shorter date range.")

        grid: list[list[str]] = []
        layout = None
        saw_text = False

        for page in pdf.pages:
            words = page.extract_words(keep_blank_chars=False, use_text_flow=False)
            if words:
                saw_text = True

            ruled = _ruled_rows(page)
            if ruled:
                grid.extend(ruled)
                continue

            rows, layout = _layout_rows(words, layout)
            grid.extend(rows)

        if not saw_text:
            raise StatementError(
                "scanned_pdf",
                "This PDF is a scanned image with no text in it. Download the statement "
                "again from net banking — the downloaded PDF has text — or use Excel.")

        return Extracted(_drop_repeated_headers(grid), "pdf", len(pdf.pages))


def _ruled_rows(page) -> list[list[str]]:
    """Rows from tables drawn with ruling lines, when the page has a real one."""
    rows: list[list[str]] = []
    for table in page.find_tables():
        lines = [[_cell_lines(page, bbox) if bbox else [] for bbox in row.cells]
                 for row in table.rows]
        width = max((len(r) for r in lines), default=0)
        # How far right each column's text reaches: the wrap width for it
        reach = [max((ln[1] for r in lines if i < len(r) for ln in r[i]), default=0.0)
                 for i in range(width)]
        cells = [[_join_wrapped(cell, reach[i]) for i, cell in enumerate(r)] for r in lines]
        if len(cells) >= 2 and width >= 3:
            rows.extend(cells)
    return rows


def _cell_lines(page, bbox) -> list[tuple[str, float, float]]:
    """A ruled cell's text lines as (text, right edge, character width)."""
    try:
        lines = page.crop(bbox).extract_text_lines(return_chars=True)
    except ValueError:            # bbox outside the page
        return []
    return [(ln["text"], ln["x1"], _char_width(ln)) for ln in lines]


def _char_width(line: dict) -> float:
    chars = line.get("chars") or []
    if chars:
        return statistics.median(c["x1"] - c["x0"] for c in chars)
    return (line["x1"] - line["x0"]) / max(1, len(line["text"]))


def _join_wrapped(lines: list[tuple[str, float, float]], reach: float) -> str:
    """
    Rejoins the lines of a wrapped cell.

    Banks wrap long unbroken strings — UPI narrations above all — mid-word:
    "UPI-SWIGGY-SWIGG" / "Y8@YBL". Joining those with a space splits "SWIGGY"
    and the transaction stops categorising as food; joining every line without
    one glues ordinary words together ("CREDIT CARDPAYMENT").

    A word-wrapping writer only cuts inside a word when that word is longer
    than the whole line, so a mid-word cut leaves a line with no spaces in it
    that runs out to the column's reach. Only such a line is joined without a
    space.
    """
    text = ""
    prev = None
    for line_text, x1, cw in lines:
        line_text = line_text.strip()
        if not line_text:
            continue
        if not text:
            text = line_text
        elif prev and " " not in prev[0] and prev[1] >= reach - 1.5 * prev[2]:
            text += line_text
        else:
            text += " " + line_text
        prev = (line_text, x1, cw)
    return text


# ── Unruled tables: lay words out under the header's columns ────────────────

_HEADER_WORDS = (
    "date", "narration", "description", "particulars", "details", "remarks",
    "amount", "debit", "withdrawal", "credit", "deposit", "balance",
    "transaction", "txn", "chq", "cheque", "ref", "value", "dr", "cr",
)
_DATE_RE = re.compile(
    r"^(\d{4}-\d{1,2}-\d{1,2}|\d{1,2}[/.\-]\d{1,2}[/.\-]\d{2,4}|\d{1,2}[\s\-/]?[A-Za-z]{3,9}[\s\-/,]*\d{2,4})$")
_AMOUNT_RE = re.compile(r"^\(?[-+]?(₹|rs\.?|inr)?\s?[\d,]*\d(\.\d{1,2})?\)?\s*(dr|cr)?\.?$", re.I)


@dataclass
class _Column:
    label: str
    x0: float
    x1: float


def _lines(words: list[dict]) -> list[list[dict]]:
    """Groups words into visual lines by their vertical position."""
    lines: list[list[dict]] = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(lines[-1][0]["top"] - w["top"]) <= 3:
            lines[-1].append(w)
        else:
            lines.append([w])
    return [sorted(line, key=lambda w: w["x0"]) for line in lines]


def _header_hits(line: list[dict]) -> int:
    # Short keywords ("dr", "cr", "ref") must be the whole word: as substrings
    # they would make "Address" and "Description" look like header cells.
    return sum(any(k == t if len(k) <= 3 else k in t for k in _HEADER_WORDS)
               for t in (w["text"].lower().strip(".:/") for w in line))


def _header_columns(line: list[dict]) -> list[_Column]:
    """Merges the header line's words into column labels ("Value", "Dt" → "Value Dt")."""
    height = statistics.median(w["bottom"] - w["top"] for w in line) or 8
    cols: list[_Column] = []
    for w in line:
        if cols and w["x0"] - cols[-1].x1 < 0.5 * height:
            cols[-1].label += " " + w["text"]
            cols[-1].x1 = w["x1"]
        else:
            cols.append(_Column(w["text"], w["x0"], w["x1"]))
    return cols


def _gutters(columns: list[_Column], lines: list[list[dict]]) -> list[float]:
    """
    Boundaries between adjacent columns, placed in the empty vertical channel
    that separates them on the page.

    Halfway between two header labels is not good enough: a narration is
    left-aligned and runs far past the midpoint towards the next label, and
    amounts are right-aligned under theirs. The channel no word crosses is
    where a reader sees the column edge, so that is what is looked for —
    using only the header and transaction lines, not footers or summaries
    that run across the page.
    """
    spans = [(w["x0"], w["x1"]) for line in lines for w in line]
    bounds = []
    for left, right in zip(columns, columns[1:]):
        lo, hi = left.x0 + 1, right.x1 - 1
        blocked = sorted((max(a, lo), min(b, hi)) for a, b in spans if b > lo and a < hi)
        # Walk the covered intervals to find the widest uncovered run in [lo, hi]
        best, best_width, cursor = None, 0.0, lo
        for a, b in blocked + [(hi, hi)]:
            if a - cursor > best_width:
                best, best_width = (cursor + a) / 2, a - cursor
            cursor = max(cursor, b)
        bounds.append(best if best is not None else (left.x1 + right.x0) / 2)
    return bounds


def _assign_words(line: list[dict], bounds: list[float]) -> list[list[dict]]:
    """Puts each word in the column whose channel-bounded span holds its centre."""
    cols: list[list[dict]] = [[] for _ in range(len(bounds) + 1)]
    for w in line:
        centre = (w["x0"] + w["x1"]) / 2
        cols[sum(1 for b in bounds if centre > b)].append(w)
    return cols


def _starts_with_date(line: list[dict]) -> bool:
    # "01/09/26" is one word; "1 Sep 2026" is three
    return bool(_DATE_RE.match(line[0]["text"])) or bool(
        _DATE_RE.match(" ".join(w["text"] for w in line[:3])))


def _layout_rows(words: list[dict], layout):
    """
    Rebuilds table rows from word positions on a page without ruling lines.

    A wrapped narration prints as extra lines with nothing in the date or
    amount columns; those are folded back into the row above instead of
    becoming rows of their own. Pages without a repeated header reuse the
    previous page's columns.
    """
    lines = _lines(words)
    header_at = next((i for i, line in enumerate(lines) if _header_hits(line) >= 3), None)

    if header_at is not None:
        columns = _header_columns(lines[header_at])
        body = lines[header_at + 1:]
        sample = [lines[header_at]] + [ln for ln in body if _starts_with_date(ln)]
        layout = (columns, _gutters(columns, sample))
    elif layout is None:
        return [], None
    else:
        body = lines

    columns, bounds = layout
    rows: list[list[str]] = [[c.label for c in columns]] if header_at is not None else []
    date_col = next((i for i, c in enumerate(columns) if "date" in c.label.lower()), 0)

    assigned = [(line, _assign_words(line, bounds)) for line in body]
    # How far right each column's text reaches: the bank's wrap width for it
    reach = [max((w["x1"] for _, cols in assigned for w in cols[i]), default=0.0)
             for i in range(len(columns))]

    prev_bottom = None
    last_line_cols = None     # word lists of the most recent line of the current row
    for line, col_words in assigned:
        cells = [" ".join(w["text"] for w in ws) for ws in col_words]
        filled = [c for c in cells if c]
        has_date = bool(_DATE_RE.match(cells[date_col]))
        has_amount = any(_AMOUNT_RE.match(c) and any(ch.isdigit() for ch in c)
                         for i, c in enumerate(cells) if i != date_col and c)
        top = min(w["top"] for w in line)
        height = max(w["bottom"] - w["top"] for w in line)
        # A wrapped narration line sits directly under the row it belongs to.
        # A page footer or summary after a gap is not part of that row, and
        # nothing continues into the date column.
        adjacent = prev_bottom is not None and top - prev_bottom < 1.2 * height
        prev_is_txn = bool(rows) and bool(_DATE_RE.match(rows[-1][date_col]))
        if (filled and not has_date and not has_amount and prev_is_txn
                and adjacent and not cells[date_col]):
            prev = rows[-1]
            for i, c in enumerate(cells):
                if not c:
                    continue
                above = last_line_cols[i] if last_line_cols else []
                if above:
                    last = above[-1]
                    cw = (last["x1"] - last["x0"]) / max(1, len(last["text"]))
                    above_text = " ".join(w["text"] for w in above)
                    joined = _join_wrapped([(above_text, last["x1"], cw), (c, 0.0, cw)], reach[i])
                    prev[i] = prev[i][: len(prev[i]) - len(above_text)] + joined
                else:
                    prev[i] = (prev[i] + " " + c).strip()
            last_line_cols = col_words
        elif filled:
            rows.append(cells)
            last_line_cols = col_words
        prev_bottom = max(w["bottom"] for w in line)
    return rows, layout


# ── Excel and friends ───────────────────────────────────────────────────────

def _extract_xlsx(data: bytes) -> list[list[str]]:
    # openpyxl parses XML with defusedxml when it is installed, which guards
    # against entity-expansion bombs in a hostile workbook.
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise StatementError("unreadable", "This Excel file could not be read.") from None
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows(values_only=True):
                rows.append([_cell(v) for v in r[:MAX_COLS]])
                if len(rows) >= MAX_ROWS:
                    break
            sheets.append(rows)
        return max(sheets, key=_filled_rows, default=[])
    finally:
        wb.close()


def _extract_xls(data: bytes) -> list[list[str]]:
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception:
        raise StatementError("unreadable", "This Excel file could not be read.") from None
    sheets = []
    for sheet in book.sheets():
        rows = []
        for r in range(min(sheet.nrows, MAX_ROWS)):
            row = []
            for c in range(min(sheet.ncols, MAX_COLS)):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate_as_datetime(cell.value, book.datemode).date().isoformat())
                        continue
                    except Exception:
                        pass
                row.append(_cell(cell.value))
            rows.append(row)
        sheets.append(rows)
    return max(sheets, key=_filled_rows, default=[])


def _extract_html(data: bytes) -> list[list[str]]:
    # Several banks' ".xls" download is an HTML table with an Excel extension.
    # Read cell text as-is: pandas.read_html would "helpfully" turn
    # "1,250.00" into 1250.0 and "0012" into 12 before we ever saw them.
    from lxml import html as lxml_html
    try:
        doc = lxml_html.fromstring(_decode(data))
    except Exception:
        raise StatementError("no_table", "No transaction table was found in this file.") from None
    tables = []
    for table in doc.iter("table"):
        rows = [[_cell(td.text_content()) for td in tr.xpath("./td|./th")][:MAX_COLS]
                for tr in table.iter("tr")]
        tables.append(rows[:MAX_ROWS])
    if not tables:
        raise StatementError("no_table", "No transaction table was found in this file.")
    return max(tables, key=_filled_rows)


def _extract_text(data: bytes) -> list[list[str]]:
    # And some are tab- or comma-separated text with an Excel extension
    text = _decode(data)
    sample = "\n".join(text.splitlines()[:40])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return [[_cell(v) for v in r[:MAX_COLS]] for r in rows[:MAX_ROWS]]


# ── Shared helpers ──────────────────────────────────────────────────────────

def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value != value:          # NaN from pandas
            return ""
        return str(int(value)) if value.is_integer() else f"{value:.2f}"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value)).strip()


def _filled_rows(rows: list[list[str]]) -> int:
    return sum(1 for r in rows if sum(1 for c in r if c) >= 3)


def _clean(grid: list[list[str]]) -> list[list[str]]:
    """Drops empty rows and trailing empty columns."""
    rows = [[_cell(c) for c in r][:MAX_COLS] for r in grid if any(_cell(c) for c in r)]
    width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
    return [(r + [""] * width)[:width] for r in rows[:MAX_ROWS]]


def _drop_repeated_headers(grid: list[list[str]]) -> list[list[str]]:
    """A statement repeats its header on every page; keep only the first."""
    seen_header = None
    out = []
    for row in grid:
        key = tuple(c.lower() for c in row)
        hits = sum(1 for c in key if any(k in c for k in _HEADER_WORDS if len(k) > 2))
        if hits >= 3:
            if seen_header == key:
                continue
            seen_header = seen_header or key
        out.append(row)
    return out
