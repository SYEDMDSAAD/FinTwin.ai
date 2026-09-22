"""
Statement extraction: PDF (ruled and unruled), Excel, and the disguised
".xls" formats Indian banks hand out, all reduced to the same grid.

Fixtures are generated in-process (fpdf2, openpyxl, xlwt) so the suite needs
no binary files and no real customer statements.
"""
import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from fpdf import FPDF
from fpdf.enums import EncryptionMethod

import app as app_module
from statements.extract import MAX_BYTES, StatementError, extract, sniff

HEADER = ["Date", "Narration", "Chq./Ref.No.", "Value Dt", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"]
TXNS = [
    ["01/09/26", "UPI-SWIGGY-SWIGGY8@YBL-YESB0YBLUPI-412345678901-PAYMENT FROM PHONEPE",
     "0000412345678901", "01/09/26", "1,250.00", "", "24,550.00"],
    ["02/09/26", "NEFT CR-ACME TECHNOLOGIES PVT LTD-SALARY SEP 2026",
     "N245261234567890", "02/09/26", "", "1,20,000.00", "1,44,550.00"],
    ["05/09/26", "CREDIT CARD PAYMENT HDFC XX1234",
     "CC0000123", "05/09/26", "12,000.00", "", "1,32,550.00"],
]


# ── fixture builders ─────────────────────────────────────────────────────────

def _preamble(pdf: FPDF):
    pdf.set_font("helvetica", size=9)
    for line in ("HDFC BANK LTD", "Account Branch : KORAMANGALA",
                 "Account No : 50100XXXXXX1234", "Statement From : 01/09/2026 To : 30/09/2026"):
        pdf.cell(0, 5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)


def ruled_pdf(pages=2) -> bytes:
    """Bordered table, long narrations wrapping inside cells, header repeated per page."""
    pdf = FPDF()
    pdf.set_font("helvetica", size=7)
    for p in range(pages):
        pdf.add_page()
        if p == 0:
            _preamble(pdf)
            pdf.set_font("helvetica", size=7)
        with pdf.table(col_widths=(14, 60, 26, 14, 20, 20, 22)) as table:
            for r in [HEADER] + TXNS:
                row = table.row()
                for c in r:
                    row.cell(c)
    return bytes(pdf.output())


# (x in mm, width, align) per column — amounts right-aligned under their header
COLS = [(10, 16, "L"), (27, 62, "L"), (90, 28, "L"), (119, 15, "L"),
        (135, 20, "R"), (156, 20, "R"), (177, 23, "R")]


def unruled_pdf(repeat_header=True, pages=2) -> bytes:
    """No ruling lines: columns exist only as x positions; narrations wrap onto extra lines."""
    pdf = FPDF()
    for p in range(pages):
        pdf.add_page()
        if p == 0:
            _preamble(pdf)
        pdf.set_font("helvetica", size=7)
        y = pdf.get_y()
        if p == 0 or repeat_header:
            for (x, w, align), text in zip(COLS, HEADER):
                pdf.set_xy(x, y)
                pdf.cell(w, 4, text, align=align)
            y += 6
        for txn in TXNS:
            # Wrap the narration inside its column, as a bank's PDF does
            lines = pdf.multi_cell(COLS[1][1], 4, txn[1], dry_run=True, output="LINES")
            for (x, w, align), text in zip(COLS, [txn[0], lines[0]] + txn[2:]):
                pdf.set_xy(x, y)
                pdf.cell(w, 4, text, align=align)
            y += 4
            for extra in lines[1:]:                   # wrapped narration lines
                pdf.set_xy(COLS[1][0], y)
                pdf.cell(COLS[1][1], 4, extra)
                y += 4
            y += 2
        pdf.set_xy(10, y + 6)
        pdf.cell(0, 4, f"Page {p + 1} of {pages}")
    return bytes(pdf.output())


def encrypted_pdf(password="SAAD0109") -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=7)
    with pdf.table() as table:
        for r in [HEADER] + TXNS:
            row = table.row()
            for c in r:
                row.cell(c)
    pdf.set_encryption(owner_password="owner", user_password=password,
                       encryption_method=EncryptionMethod.AES_256)
    return bytes(pdf.output())


def scanned_pdf() -> bytes:
    from PIL import Image
    img = io.BytesIO()
    Image.new("RGB", (400, 200), "white").save(img, format="PNG")
    pdf = FPDF()
    pdf.add_page()
    pdf.image(io.BytesIO(img.getvalue()), x=10, y=10, w=100)
    return bytes(pdf.output())


def xlsx_bytes() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Account Statement"])
    ws.append(["Account No", "XXXX1234"])
    ws.append([])
    ws.append(["Transaction Date", "Description", "Debit", "Credit", "Balance"])
    ws.append([date(2026, 9, 1), "UPI/SWIGGY", 1250.0, None, 24550.0])
    ws.append([date(2026, 9, 2), "NEFT CR ACME PAYROLL", None, 120000.0, 144550.5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xls_bytes() -> bytes:
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Statement")
    date_style = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for c, h in enumerate(["Txn Date", "Description", "Debit", "Credit", "Balance"]):
        ws.write(0, c, h)
    ws.write(1, 0, date(2026, 9, 1), date_style)
    ws.write(1, 1, "UPI/SWIGGY")
    ws.write(1, 2, 1250.0)
    ws.write(1, 4, 24550.0)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HTML_XLS = (b"<html><body><table>"
            b"<tr><td>Account No</td><td>XXXX1234</td></tr>"
            b"<tr><th>Date</th><th>Particulars</th><th>Debit</th><th>Credit</th><th>Balance</th></tr>"
            b"<tr><td>01/09/2026</td><td>UPI/SWIGGY</td><td>1,250.00</td><td></td><td>24,550.00</td></tr>"
            b"</table></body></html>")

TAB_XLS = (b"Date\tNarration\tDebit\tCredit\tBalance\r\n"
           b"01/09/2026\tUPI/SWIGGY\t1,250.00\t\t24,550.00\r\n")


def _row_with(grid, text):
    return next(r for r in grid if any(text in c for c in r))


# ── type sniffing ────────────────────────────────────────────────────────────

def test_sniff_goes_by_content_not_extension():
    assert sniff(ruled_pdf(1)) == "pdf"
    assert sniff(xlsx_bytes()) == "xlsx"
    assert sniff(xls_bytes()) == "xls"
    assert sniff(HTML_XLS) == "html"
    assert sniff(TAB_XLS) == "text"
    assert sniff(bytes(range(256)) * 4) == "unknown"


# ── PDF ──────────────────────────────────────────────────────────────────────

def test_ruled_pdf_keeps_wrapped_narrations_whole_and_one_header():
    result = extract(ruled_pdf(pages=2))
    grid = result.grid

    assert result.format == "pdf" and result.pages == 2
    headers = [r for r in grid if r[0] == "Date"]
    assert len(headers) == 1, "the header repeated on page 2 must be dropped"
    swiggy = _row_with(grid, "SWIGGY")
    # Wrapped inside its cell — mid-word through the UPI string, at a space
    # before "PHONEPE" — and read back exactly as written
    assert swiggy[1] == TXNS[0][1]
    assert swiggy[4] == "1,250.00" and swiggy[6] == "24,550.00"
    # both pages' transactions are present
    assert sum(1 for r in grid if "ACME" in " ".join(r)) == 2


@pytest.mark.parametrize("repeat_header", [True, False])
def test_unruled_pdf_is_rebuilt_from_word_positions(repeat_header):
    grid = extract(unruled_pdf(repeat_header=repeat_header)).grid

    header = grid[0]
    assert header[:2] == ["Date", "Narration"]
    assert "Value Dt" in header and "Withdrawal Amt." in header

    # Page footers ("Page 1 of 2") survive as rows without a date; the
    # importer downstream skips undated rows.
    rows = [r for r in grid[1:] if r[0][:2].isdigit() and "/" in r[0]]
    assert len(rows) == 6, "3 transactions on each of 2 pages, no continuation rows"
    salary = _row_with(rows, "ACME")
    # the wrapped lines are folded back into the same narration
    assert salary[1] == TXNS[1][1]
    col = header.index
    assert salary[col("Deposit Amt.")] == "1,20,000.00"
    assert salary[col("Withdrawal Amt.")] == ""
    assert salary[col("Closing Balance")] == "1,44,550.00"
    card = _row_with(rows, "CREDIT CARD")
    assert card[col("Withdrawal Amt.")] == "12,000.00"


def _long_unruled_statement(n=90, seed=7, repeat_header=True):
    """A realistic multi-page statement: wrapped narrations, page breaks, footers."""
    import random
    rnd = random.Random(seed)
    merchants = [
        "UPI-ZOMATO-ZOMATO@HDFCBANK-HDFC0MERUPI-512345678901-PAYMENT",
        "ATW-512345XXXXXX1234-S1ANMU12-BANGALORE", "POS 512345XXXXXX1234 AMAZON PAY INDIA",
        "NEFT CR-ICICI0000123-RAHUL KUMAR-RENT SHARE", "ACH D- TP ACH ZERODHA BROKING LTD-1234567",
        "IMPS-612345678901-MOM-SBIN0001234-XXXXXXX5678-GIFT", "CRED CLUB PAYMENT CREDIT CARD BILL",
        "NETFLIX.COM",
    ]
    money = "{:,.2f}".format
    txns, bal = [], 50_000.0
    for i in range(n):
        d = f"{i % 28 + 1:02d}/09/26"
        amt = round(rnd.uniform(1, 250_000), 2)
        credit = rnd.random() < 0.25
        bal += amt if credit else -amt
        txns.append([d, rnd.choice(merchants), str(rnd.randint(10**11, 10**12)), d,
                     "" if credit else money(amt), money(amt) if credit else "", money(bal)])

    pdf = FPDF()
    pdf.set_auto_page_break(False)

    def header(y):
        for (x, w, align), text in zip(COLS, HEADER):
            pdf.set_xy(x, y)
            pdf.cell(w, 4, text, align=align)
        return y + 6

    pdf.add_page()
    _preamble(pdf)
    pdf.set_font("helvetica", size=7)
    y, page = header(pdf.get_y()), 1
    for t in txns:
        lines = pdf.multi_cell(COLS[1][1], 4, t[1], dry_run=True, output="LINES")
        if y + 4 * len(lines) > 275:
            pdf.set_xy(10, 285)
            pdf.cell(0, 4, f"Page {page}")
            pdf.add_page()
            pdf.set_font("helvetica", size=7)
            page += 1
            y = header(15) if repeat_header else 15
        for (x, w, align), text in zip(COLS, [t[0], lines[0]] + t[2:]):
            pdf.set_xy(x, y)
            pdf.cell(w, 4, text, align=align)
        y += 4
        for extra in lines[1:]:
            pdf.set_xy(COLS[1][0], y)
            pdf.cell(COLS[1][1], 4, extra)
            y += 4
        y += 2
    return bytes(pdf.output()), txns


@pytest.mark.parametrize("repeat_header", [True, False])
def test_long_unruled_statement_puts_every_amount_in_its_column(repeat_header):
    data, txns = _long_unruled_statement(repeat_header=repeat_header)
    result = extract(data)
    assert result.pages >= 3

    rows = [r for r in result.grid[1:] if "/" in r[0]]
    assert len(rows) == len(txns)
    for got, want in zip(rows, txns):
        # date, withdrawal, deposit, balance exact; footers never leak into a row
        assert [got[0], got[4], got[5], got[6]] == [want[0], want[4], want[5], want[6]]
        assert got[1] == want[1]


def test_narrow_cells_rejoin_mid_word_wraps_without_gluing_words():
    # The encrypted fixture's auto-sized table is narrow: the UPI string is
    # cut mid-word several times, ordinary narrations wrap at spaces
    grid = extract(encrypted_pdf("pw"), password="pw").grid
    narrations = [r[1] for r in grid if "/" in r[0]]
    # "SWIGG Y8" would stop this categorising as food; "CARDPAYMENT" would
    # stop the bill-payment rule matching
    assert narrations == [t[1] for t in TXNS]


def upi_app_pdf() -> bytes:
    """PhonePe's layout: month-first dates, the time on the line below, a
    Type column, and an amount that sometimes wraps ("INR" / "12345.00")."""
    cols = [(10, 30, "L"), (42, 90, "L"), (134, 22, "L"), (158, 40, "L")]
    txns = [
        ("Jun 24, 2026", "05:05 PM", "Received from ******1317", "Credit", "INR 1000.00", None),
        ("Jun 24, 2026", "05:12 PM", "Paid to Apple Services", "Debit", "INR 39.00", None),
        ("Jul 02, 2026", "11:40 AM", "Received from ******4411", "Credit", "INR", "25000.00"),
        ("Jul 03, 2026", "09:15 PM", "Paid - Mobile Recharge", "Debit", "INR 33.00", None),
    ]
    pdf = FPDF()
    pdf.set_auto_page_break(False)
    pdf.add_page()
    pdf.set_font("helvetica", size=8)
    y = 20
    for (x, w, a), text in zip(cols, ["Date", "Transaction Details", "Type", "Amount"]):
        pdf.set_xy(x, y); pdf.cell(w, 4, text, align=a)
    y += 8
    for date_, time_, details, typ, amount, wrapped in txns:
        for (x, w, a), text in zip(cols, [date_, details, typ, amount]):
            pdf.set_xy(x, y); pdf.cell(w, 4, text, align=a)
        y += 4
        for (x, w, a), text in zip(cols, [time_, "Transaction ID : T2606241705189563531737", "", wrapped or ""]):
            if text:
                pdf.set_xy(x, y); pdf.cell(w, 4, text, align=a)
        y += 4
        pdf.set_xy(cols[1][0], y); pdf.cell(cols[1][1], 4, "UTR No : 450474129626")
        y += 8
    pdf.set_xy(10, 280)
    pdf.cell(0, 4, "This is a system generated statement. For any queries contact us at https://support.phonepe.com")
    return bytes(pdf.output())


def test_upi_app_layout_keeps_every_transaction_and_its_amount():
    grid = extract(upi_app_pdf()).grid
    rows = [r for r in grid if r[0][:3] in ("Jun", "Jul")]

    assert [r[0] for r in rows] == ["Jun 24, 2026", "Jun 24, 2026", "Jul 02, 2026", "Jul 03, 2026"]
    assert [r[2] for r in rows] == ["Credit", "Debit", "Credit", "Debit"]
    # the wrapped amount is rejoined — not lost as a row with no date
    assert rows[2][3] == "INR 25000.00"
    # the time and Transaction ID lines are not folded into the payee
    assert rows[1][1] == "Paid to Apple Services"


def test_password_protected_pdf_asks_then_opens():
    data = encrypted_pdf("SAAD0109")

    with pytest.raises(StatementError) as missing:
        extract(data)
    assert missing.value.code == "password_required"

    with pytest.raises(StatementError) as wrong:
        extract(data, password="nope")
    assert wrong.value.code == "password_incorrect"

    grid = extract(data, password="SAAD0109").grid
    assert _row_with(grid, "SWIGGY")[4] == "1,250.00"


def test_scanned_pdf_is_reported_not_returned_empty():
    with pytest.raises(StatementError) as err:
        extract(scanned_pdf())
    assert err.value.code == "scanned_pdf"


# ── Excel and disguised Excel ────────────────────────────────────────────────

def test_xlsx_dates_and_numbers_come_back_as_plain_strings():
    grid = extract(xlsx_bytes()).grid
    header = _row_with(grid, "Transaction Date")
    assert header == ["Transaction Date", "Description", "Debit", "Credit", "Balance"]
    assert _row_with(grid, "SWIGGY") == ["2026-09-01", "UPI/SWIGGY", "1250", "", "24550"]
    assert _row_with(grid, "ACME")[4] == "144550.50"
    assert grid[0][0] == "Account Statement", "preamble is kept for the header finder downstream"


def test_legacy_xls_reads_real_date_cells():
    grid = extract(xls_bytes()).grid
    assert grid[1][:3] == ["2026-09-01", "UPI/SWIGGY", "1250"]


def test_html_saved_as_xls_is_read_as_a_table():
    result = extract(HTML_XLS)
    assert result.format == "html"
    assert _row_with(result.grid, "SWIGGY") == ["01/09/2026", "UPI/SWIGGY", "1,250.00", "", "24,550.00"]


def test_tab_text_saved_as_xls_is_read_as_a_table():
    result = extract(TAB_XLS)
    assert result.format == "text"
    assert result.grid[1] == ["01/09/2026", "UPI/SWIGGY", "1,250.00", "", "24,550.00"]


def test_rejects_empty_oversized_and_unknown_files():
    for data, code in [(b"", "empty"), (b"x" * (MAX_BYTES + 1), "too_large"),
                       (bytes(range(256)) * 4, "unsupported")]:
        with pytest.raises(StatementError) as err:
            extract(data)
        assert err.value.code == code


# ── HTTP route ───────────────────────────────────────────────────────────────

client = TestClient(app_module.app)
KEY = {"x-internal-key": "test-internal-key"}


def test_route_requires_the_internal_key():
    r = client.post("/statements/extract", files={"file": ("s.pdf", ruled_pdf(1))})
    assert r.status_code == 403


def test_route_returns_the_grid():
    r = client.post("/statements/extract", headers=KEY,
                    files={"file": ("s.xlsx", xlsx_bytes())})
    assert r.status_code == 200
    body = r.json()
    assert body["format"] == "xlsx"
    assert ["2026-09-01", "UPI/SWIGGY", "1250", "", "24550"] in body["grid"]


def test_route_reports_password_errors_with_a_code():
    data = encrypted_pdf("pw")
    r = client.post("/statements/extract", headers=KEY, files={"file": ("s.pdf", data)})
    assert r.status_code == 422
    assert r.json()["detail"]["code"] == "password_required"

    r = client.post("/statements/extract", headers=KEY,
                    files={"file": ("s.pdf", data)}, data={"password": "pw"})
    assert r.status_code == 200


def test_a_statement_too_long_to_read_is_refused_not_cut_short(monkeypatch):
    import statements.extract as ex
    monkeypatch.setattr(ex, "MAX_ROWS", 3)
    csv_text = "Date,Details,Amount\n" + "\n".join(f"0{i}/09/2026,Shop {i},-{i}00" for i in range(1, 6))
    with pytest.raises(StatementError) as e:
        extract(csv_text.encode(), None)
    assert e.value.code == "too_many_rows"
    assert "import it in parts" in e.value.message


def test_ruled_cells_are_read_from_one_pass_over_the_page():
    fpdf = pytest.importorskip("fpdf")
    pdf = fpdf.FPDF()
    pdf.set_font("Helvetica", size=8)
    for _ in range(3):
        pdf.add_page()
        for row in [("Date", "Details", "Debit", "Balance"),
                    ("01/09/2026", "UPI/SWIGGY/ORDER", "250.00", "9750.00"),
                    ("02/09/2026", "UPI/ZEPTO/ORDER", "499.00", "9251.00")]:
            for text, w in zip(row, (25, 70, 25, 25)):
                pdf.cell(w, 7, text, border=1)
            pdf.ln()
    grid = extract(bytes(pdf.output()), None).grid
    # the header repeats on every page and is kept once; each cell holds only its own text
    assert grid[0] == ["Date", "Details", "Debit", "Balance"]
    assert grid[1:] == [["01/09/2026", "UPI/SWIGGY/ORDER", "250.00", "9750.00"],
                        ["02/09/2026", "UPI/ZEPTO/ORDER", "499.00", "9251.00"]] * 3
